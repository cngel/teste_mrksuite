import io
import os
import uuid
from datetime import date

import psycopg2
from fastapi import Depends, FastAPI, File, HTTPException, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer
from jose import JWTError, jwt
from psycopg2.extras import RealDictCursor
from pydantic import BaseModel

if __package__:
    from .core.db import (
        EMPLOYEE_SUBFOLDERS,
        categories_for_department,
        ensure_default_tags,
        ensure_top_level_folders,
        get_connection,
        get_or_create_top_folder,
        init_db,
        resolve_folder_permission,
    )
    from .core.tools import BUCKET_NAME, delete_object, ensure_bucket, get_minio_client, get_presigned_url, upload_bytes
else:
    from core.db import (
    EMPLOYEE_SUBFOLDERS,
    categories_for_department,
    ensure_default_tags,
    ensure_top_level_folders,
    get_connection,
    get_or_create_top_folder,
    init_db,
    resolve_folder_permission,
    )
    from core.tools import BUCKET_NAME, delete_object, ensure_bucket, get_minio_client, get_presigned_url, upload_bytes

app = FastAPI(title="Documents Service", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=os.environ.get(
        "CORS_ORIGINS",
        "http://localhost,http://localhost:80,http://127.0.0.1,http://localhost:5002",
    ).split(","),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

JWT_SECRET    = os.environ["JWT_SECRET"]
JWT_ALGORITHM = "HS256"
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="http://localhost:5000/login")


@app.on_event("startup")
def on_startup():
    init_db()
    ensure_bucket()


# ---------------------------------------------------------------------------
# Verificação JWT (partilha blocklist com auth-service via mesma DB)
# ---------------------------------------------------------------------------

def get_current_claims(token: str = Depends(oauth2_scheme)) -> dict:
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except JWTError:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Token inválido")

    if payload.get("type") != "access":
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Token de acesso inválido")

    jti = payload.get("jti")
    if jti:
        with get_connection() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT jti FROM jwt_blocklist WHERE jti = %s AND expires_at > NOW()",
                    (jti,),
                )
                if cur.fetchone():
                    raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Token revogado")

    return payload


def get_current_user(claims: dict = Depends(get_current_claims)) -> str:
    return claims["sub"]


def get_current_company(claims: dict = Depends(get_current_claims)) -> int:
    """Isolamento entre empresas: toda a leitura/escrita de pastas e documentos tem
    de ser filtrada por company_id — nunca basta um JWT válido, seja lá de que
    empresa for."""
    company_id = claims.get("company_id")
    if company_id is None:
        raise HTTPException(status_code=403, detail="Utilizador sem empresa associada")
    return company_id


def resolve_user_name(cur, user_id: str) -> str:
    cur.execute("SELECT nome FROM usuarios WHERE id = %s", (user_id,))
    row = cur.fetchone()
    return row["nome"] if row else "Utilizador"


def require_permission(cur, user_id: str, folder_id: int | None, action: str):
    perm = resolve_folder_permission(cur, user_id, folder_id)
    if not perm[f"can_{action}"]:
        raise HTTPException(status_code=403, detail="Sem permissão para esta pasta")


def notify(cur, user_id: str | None, notif_type: str, message: str, company_id: int, document_id: int | None = None):
    if not user_id:
        return
    cur.execute(
        "INSERT INTO doc_notifications (user_id, type, document_id, message, company_id) VALUES (%s, %s, %s, %s, %s)",
        (user_id, notif_type, document_id, message, company_id),
    )


# ---------------------------------------------------------------------------
# MODELS
# ---------------------------------------------------------------------------

class FolderCreate(BaseModel):
    name: str
    parent_id: int | None = None


class FolderUpdate(BaseModel):
    name: str


class EnsureEmployeeFolder(BaseModel):
    employee_name: str


class EnsureDepartmentFolder(BaseModel):
    department_name: str


class DocumentUpdate(BaseModel):
    name: str | None = None
    folder_id: int | None = None
    category: str | None = None
    department_id: int | None = None
    contact_id: int | None = None
    project_id: int | None = None
    expiry_date: str | None = None


class DocumentTagsUpdate(BaseModel):
    tag_ids: list[int]


class SignatureUpdate(BaseModel):
    signature_status: str


class TagCreate(BaseModel):
    name: str
    color: str = "#6B7280"


class WorkflowTransition(BaseModel):
    notes: str | None = None


class PermissionCreate(BaseModel):
    scope_type: str
    scope_value: str | None = None
    folder_id: int
    can_view: bool = True
    can_edit: bool = False
    can_delete: bool = False


class TemplateUse(BaseModel):
    folder_id: int
    fields: dict[str, str] = {}
    document_name: str | None = None


# ---------------------------------------------------------------------------
# Folders
# ---------------------------------------------------------------------------

@app.get("/folders")
def list_folders(user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            ensure_top_level_folders(cur, company_id)
            conn.commit()
            cur.execute("SELECT * FROM folders WHERE company_id = %s ORDER BY is_system DESC, name", (company_id,))
            return cur.fetchall()


@app.post("/folders", status_code=201)
def create_folder(body: FolderCreate, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            if body.parent_id is not None:
                cur.execute("SELECT 1 FROM folders WHERE id = %s AND company_id = %s", (body.parent_id, company_id))
                if not cur.fetchone():
                    raise HTTPException(status_code=404, detail="Pasta-mãe não encontrada")
            cur.execute(
                "INSERT INTO folders (name, parent_id, kind, company_id) VALUES (%s, %s, 'custom', %s) RETURNING *",
                (body.name, body.parent_id, company_id),
            )
            conn.commit()
            return cur.fetchone()


@app.put("/folders/{folder_id}")
def update_folder(folder_id: int, body: FolderUpdate, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT * FROM folders WHERE id = %s AND company_id = %s", (folder_id, company_id))
            folder = cur.fetchone()
            if not folder:
                raise HTTPException(status_code=404, detail="Pasta não encontrada")
            if folder["is_system"]:
                raise HTTPException(status_code=403, detail="Não é possível renomear uma pasta do sistema")
            require_permission(cur, user_id, folder["parent_id"], "edit")

            cur.execute(
                "UPDATE folders SET name=%s WHERE id=%s AND company_id=%s RETURNING *",
                (body.name, folder_id, company_id),
            )
            conn.commit()
            return cur.fetchone()


@app.delete("/folders/{folder_id}", status_code=204)
def delete_folder(folder_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT * FROM folders WHERE id = %s AND company_id = %s", (folder_id, company_id))
            folder = cur.fetchone()
            if not folder:
                raise HTTPException(status_code=404, detail="Pasta não encontrada")
            if folder["is_system"]:
                raise HTTPException(status_code=403, detail="Não é possível apagar uma pasta do sistema")
            require_permission(cur, user_id, folder["parent_id"], "delete")

            cur.execute("SELECT 1 FROM folders WHERE parent_id = %s LIMIT 1", (folder_id,))
            if cur.fetchone():
                raise HTTPException(status_code=409, detail="A pasta tem subpastas")
            cur.execute("SELECT 1 FROM documents WHERE folder_id = %s LIMIT 1", (folder_id,))
            if cur.fetchone():
                raise HTTPException(status_code=409, detail="A pasta tem documentos")

            cur.execute("DELETE FROM folders WHERE id = %s AND company_id = %s", (folder_id, company_id))
            conn.commit()


@app.post("/employees/{employee_id}/ensure-folder")
def ensure_employee_folder(employee_id: int, body: EnsureEmployeeFolder, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            colaboradores_id = get_or_create_top_folder(cur, "Colaboradores", company_id)

            cur.execute(
                "SELECT * FROM folders WHERE kind = 'employee' AND employee_id = %s AND company_id = %s",
                (employee_id, company_id),
            )
            emp_folder = cur.fetchone()
            if not emp_folder:
                cur.execute(
                    """INSERT INTO folders (name, parent_id, kind, employee_id, company_id)
                       VALUES (%s, %s, 'employee', %s, %s) RETURNING *""",
                    (body.employee_name, colaboradores_id, employee_id, company_id),
                )
                emp_folder = cur.fetchone()

            cur.execute(
                "SELECT name FROM folders WHERE parent_id = %s AND kind = 'employee_sub'",
                (emp_folder["id"],),
            )
            existing_sub = {row["name"] for row in cur.fetchall()}
            for sub_name in EMPLOYEE_SUBFOLDERS:
                if sub_name not in existing_sub:
                    cur.execute(
                        "INSERT INTO folders (name, parent_id, kind, company_id) VALUES (%s, %s, 'employee_sub', %s)",
                        (sub_name, emp_folder["id"], company_id),
                    )

            conn.commit()

            cur.execute("SELECT * FROM folders WHERE parent_id = %s ORDER BY name", (emp_folder["id"],))
            subfolders = cur.fetchall()

    return {"folder": emp_folder, "subfolders": subfolders}


@app.post("/departments/{department_id}/ensure-folder")
def ensure_department_folder(department_id: int, body: EnsureDepartmentFolder, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            departamentos_id = get_or_create_top_folder(cur, "Departamentos", company_id)

            cur.execute(
                "SELECT * FROM folders WHERE kind = 'department' AND department_id = %s AND company_id = %s",
                (department_id, company_id),
            )
            dept_folder = cur.fetchone()
            if not dept_folder:
                cur.execute(
                    """INSERT INTO folders (name, parent_id, kind, department_id, company_id)
                       VALUES (%s, %s, 'department', %s, %s) RETURNING *""",
                    (body.department_name, departamentos_id, department_id, company_id),
                )
                dept_folder = cur.fetchone()

            cur.execute(
                "SELECT id, name FROM folders WHERE parent_id = %s AND kind = 'dept_category'",
                (dept_folder["id"],),
            )
            existing_categories = {row["name"]: row["id"] for row in cur.fetchall()}

            for category_name, subcategories in categories_for_department(body.department_name):
                category_id = existing_categories.get(category_name)
                if not category_id:
                    cur.execute(
                        "INSERT INTO folders (name, parent_id, kind, company_id) VALUES (%s, %s, 'dept_category', %s) RETURNING id",
                        (category_name, dept_folder["id"], company_id),
                    )
                    category_id = cur.fetchone()["id"]

                cur.execute(
                    "SELECT name FROM folders WHERE parent_id = %s AND kind = 'dept_subcategory'",
                    (category_id,),
                )
                existing_sub = {row["name"] for row in cur.fetchall()}
                for sub_name in subcategories:
                    if sub_name not in existing_sub:
                        cur.execute(
                            "INSERT INTO folders (name, parent_id, kind, company_id) VALUES (%s, %s, 'dept_subcategory', %s)",
                            (sub_name, category_id, company_id),
                        )

            cur.execute(
                """SELECT 1 FROM document_permissions
                   WHERE folder_id = %s AND scope_type = 'department' AND scope_value = %s""",
                (dept_folder["id"], str(department_id)),
            )
            if not cur.fetchone():
                cur.execute(
                    """INSERT INTO document_permissions
                       (scope_type, scope_value, folder_id, can_view, can_edit, can_delete, company_id)
                       VALUES ('department', %s, %s, true, true, false, %s)""",
                    (str(department_id), dept_folder["id"], company_id),
                )

            conn.commit()

    return {"folder": dept_folder}


# ---------------------------------------------------------------------------
# Documents
# ---------------------------------------------------------------------------

def attach_tags_and_url(cur, doc):
    cur.execute("""
        SELECT t.id, t.name, t.color FROM document_tags dt
        JOIN doc_tags t ON t.id = dt.tag_id WHERE dt.document_id = %s
    """, (doc["id"],))
    doc["tags"] = cur.fetchall()
    try:
        doc["url"] = get_presigned_url(doc["object_name"])
    except Exception:
        doc["url"] = None
    return doc


def purge_expired_trash(cur, company_id: int):
    cur.execute("""
        SELECT id, object_name FROM documents
        WHERE company_id = %s AND deleted_at IS NOT NULL AND deleted_at < NOW() - INTERVAL '30 days'
    """, (company_id,))
    expired = cur.fetchall()
    for doc in expired:
        delete_object(doc["object_name"])
        cur.execute("DELETE FROM documents WHERE id = %s", (doc["id"],))


@app.get("/documents")
def list_documents(
    folder_id: int | None = None,
    search: str | None = None,
    tag_id: int | None = None,
    category: str | None = None,
    department_id: int | None = None,
    status_filter: str | None = None,
    favorite: bool | None = None,
    trashed: bool = False,
    expiring_before: str | None = None,
    user_id: str = Depends(get_current_user),
    company_id: int = Depends(get_current_company),
):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            if trashed:
                purge_expired_trash(cur, company_id)
                conn.commit()

            clauses = ["company_id = %s", "deleted_at IS NOT NULL" if trashed else "deleted_at IS NULL"]
            params = [company_id]

            if folder_id is not None:
                clauses.append("folder_id = %s")
                params.append(folder_id)
            if search:
                clauses.append("name ILIKE %s")
                params.append(f"%{search}%")
            if category:
                clauses.append("category = %s")
                params.append(category)
            if department_id is not None:
                clauses.append("department_id = %s")
                params.append(department_id)
            if status_filter:
                clauses.append("status = %s")
                params.append(status_filter)
            if favorite:
                clauses.append("is_favorite = true")
            if expiring_before:
                clauses.append("expiry_date IS NOT NULL AND expiry_date <= %s")
                params.append(expiring_before)
            if tag_id is not None:
                clauses.append("id IN (SELECT document_id FROM document_tags WHERE tag_id = %s)")
                params.append(tag_id)

            where_sql = " AND ".join(clauses)
            cur.execute(
                f"SELECT * FROM documents WHERE {where_sql} ORDER BY updated_at DESC",
                params,
            )
            docs = cur.fetchall()

            perm_cache: dict[int, bool] = {}
            visible = []
            for doc in docs:
                fid = doc["folder_id"]
                if fid not in perm_cache:
                    perm_cache[fid] = resolve_folder_permission(cur, user_id, fid)["can_view"]
                if perm_cache[fid]:
                    visible.append(doc)

            return [attach_tags_and_url(cur, d) for d in visible]


@app.post("/documents", status_code=201)
def upload_document(
    folder_id: int,
    category: str | None = None,
    department_id: int | None = None,
    contact_id: int | None = None,
    project_id: int | None = None,
    expiry_date: str | None = None,
    file: UploadFile = File(...),
    user_id: str = Depends(get_current_user),
    company_id: int = Depends(get_current_company),
):
    data = file.file.read()
    object_name = f"{folder_id}/{uuid.uuid4()}_{file.filename}"

    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT 1 FROM folders WHERE id = %s AND company_id = %s", (folder_id, company_id))
            if not cur.fetchone():
                raise HTTPException(status_code=404, detail="Pasta não encontrada")
            require_permission(cur, user_id, folder_id, "edit")
            owner_name = resolve_user_name(cur, user_id)
            upload_bytes(data, object_name, file.content_type or "application/octet-stream")
            cur.execute(
                """INSERT INTO documents
                       (folder_id, name, object_name, content_type, size_bytes,
                        category, department_id, contact_id, project_id, owner_id, owner_name, expiry_date,
                        workflow_state, company_id)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'enviado', %s)
                   RETURNING *""",
                (
                    folder_id, file.filename, object_name, file.content_type, len(data),
                    category, department_id, contact_id, project_id, user_id, owner_name, expiry_date,
                    company_id,
                ),
            )
            doc = cur.fetchone()
            cur.execute(
                """INSERT INTO document_versions
                       (document_id, version_number, object_name, content_type, size_bytes, uploaded_by, uploaded_by_name, company_id)
                   VALUES (%s, 1, %s, %s, %s, %s, %s, %s)""",
                (doc["id"], object_name, file.content_type, len(data), user_id, owner_name, company_id),
            )
            notify(cur, user_id, "upload_concluido", f"Upload concluído: {doc['name']}", company_id, doc["id"])
            conn.commit()
            return attach_tags_and_url(cur, doc)


@app.get("/documents/{document_id}")
def get_document(document_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT * FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, doc["folder_id"], "view")
            return attach_tags_and_url(cur, doc)


@app.put("/documents/{document_id}")
def update_document(document_id: int, body: DocumentUpdate, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT * FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, doc["folder_id"], "edit")

            if body.folder_id is not None:
                cur.execute("SELECT 1 FROM folders WHERE id = %s AND company_id = %s", (body.folder_id, company_id))
                if not cur.fetchone():
                    raise HTTPException(status_code=404, detail="Pasta não encontrada")

            # contact_id/project_id usam model_fields_set (em vez de "is not None") para
            # distinguir "campo não enviado" de "enviado como null para limpar a ligação".
            fields_set = body.model_fields_set
            cur.execute(
                """UPDATE documents SET
                       name=%s, folder_id=%s, category=%s, department_id=%s,
                       contact_id=%s, project_id=%s, expiry_date=%s, updated_at=NOW()
                   WHERE id=%s RETURNING *""",
                (
                    body.name if body.name is not None else doc["name"],
                    body.folder_id if body.folder_id is not None else doc["folder_id"],
                    body.category if body.category is not None else doc["category"],
                    body.department_id if body.department_id is not None else doc["department_id"],
                    body.contact_id if "contact_id" in fields_set else doc["contact_id"],
                    body.project_id if "project_id" in fields_set else doc["project_id"],
                    body.expiry_date if body.expiry_date is not None else doc["expiry_date"],
                    document_id,
                ),
            )
            conn.commit()
            return attach_tags_and_url(cur, cur.fetchone())


@app.put("/documents/{document_id}/tags")
def set_document_tags(document_id: int, body: DocumentTagsUpdate, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT folder_id FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, doc["folder_id"], "edit")

            cur.execute("DELETE FROM document_tags WHERE document_id = %s", (document_id,))
            for tag_id in body.tag_ids:
                cur.execute(
                    "INSERT INTO document_tags (document_id, tag_id) VALUES (%s, %s)",
                    (document_id, tag_id),
                )
            conn.commit()
            cur.execute("""
                SELECT t.id, t.name, t.color FROM document_tags dt
                JOIN doc_tags t ON t.id = dt.tag_id WHERE dt.document_id = %s
            """, (document_id,))
            return cur.fetchall()


@app.post("/documents/{document_id}/favorite")
def toggle_favorite(document_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(
                "UPDATE documents SET is_favorite = NOT is_favorite, updated_at=NOW() WHERE id=%s AND company_id=%s RETURNING *",
                (document_id, company_id),
            )
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            conn.commit()
            return attach_tags_and_url(cur, doc)


@app.post("/documents/{document_id}/archive")
def toggle_archive(document_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT folder_id FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            existing = cur.fetchone()
            if not existing:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, existing["folder_id"], "edit")

            cur.execute(
                """UPDATE documents SET
                       status = CASE WHEN status = 'arquivado' THEN 'ativo' ELSE 'arquivado' END,
                       updated_at = NOW()
                   WHERE id=%s RETURNING *""",
                (document_id,),
            )
            doc = cur.fetchone()
            conn.commit()
            return attach_tags_and_url(cur, doc)


@app.patch("/documents/{document_id}/signature")
def update_signature(document_id: int, body: SignatureUpdate, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(
                "SELECT folder_id, owner_id, name FROM documents WHERE id = %s AND company_id = %s",
                (document_id, company_id),
            )
            existing = cur.fetchone()
            if not existing:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, existing["folder_id"], "edit")

            signer_name = resolve_user_name(cur, user_id)
            cur.execute(
                """UPDATE documents SET
                       signature_status=%s, signed_by=%s, signed_at=NOW(), updated_at=NOW()
                   WHERE id=%s AND company_id=%s RETURNING *""",
                (body.signature_status, signer_name, document_id, company_id),
            )
            doc = cur.fetchone()

            if body.signature_status == "assinado":
                notify(cur, existing["owner_id"], "documento_assinado", f"Documento assinado: {existing['name']}", company_id, document_id)
            elif body.signature_status == "rejeitado":
                notify(cur, existing["owner_id"], "documento_rejeitado", f"Assinatura rejeitada: {existing['name']}", company_id, document_id)
            elif body.signature_status in ("pendente", "em_assinatura"):
                notify(cur, existing["owner_id"], "assinatura_pendente", f"Assinatura pendente: {existing['name']}", company_id, document_id)

            conn.commit()
            return attach_tags_and_url(cur, doc)


@app.delete("/documents/{document_id}", status_code=204)
def trash_document(document_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT folder_id FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            existing = cur.fetchone()
            if not existing:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, existing["folder_id"], "delete")

            cur.execute("UPDATE documents SET deleted_at = NOW() WHERE id = %s", (document_id,))
            conn.commit()


@app.post("/documents/{document_id}/restore")
def restore_document(document_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT folder_id FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            existing = cur.fetchone()
            if not existing:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, existing["folder_id"], "edit")

            cur.execute(
                "UPDATE documents SET deleted_at = NULL, updated_at = NOW() WHERE id=%s RETURNING *",
                (document_id,),
            )
            doc = cur.fetchone()
            conn.commit()
            return attach_tags_and_url(cur, doc)


@app.delete("/documents/{document_id}/permanent", status_code=204)
def delete_document_permanent(document_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT * FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, doc["folder_id"], "delete")
            if not doc["deleted_at"]:
                raise HTTPException(status_code=400, detail="Só é possível apagar definitivamente a partir da lixeira")

            delete_object(doc["object_name"])
            cur.execute("DELETE FROM documents WHERE id = %s", (document_id,))
            conn.commit()


# ---------------------------------------------------------------------------
# Versions
# ---------------------------------------------------------------------------

@app.get("/documents/{document_id}/versions")
def list_versions(document_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(
                "SELECT folder_id, current_version FROM documents WHERE id = %s AND company_id = %s",
                (document_id, company_id),
            )
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, doc["folder_id"], "view")

            cur.execute(
                "SELECT * FROM document_versions WHERE document_id = %s ORDER BY version_number DESC",
                (document_id,),
            )
            versions = cur.fetchall()
            for v in versions:
                v["is_current"] = v["version_number"] == doc["current_version"]
            return versions


@app.post("/documents/{document_id}/versions", status_code=201)
def upload_version(
    document_id: int,
    notes: str | None = None,
    file: UploadFile = File(...),
    user_id: str = Depends(get_current_user),
    company_id: int = Depends(get_current_company),
):
    data = file.file.read()

    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT * FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, doc["folder_id"], "edit")

            object_name = f"{doc['folder_id']}/{uuid.uuid4()}_{file.filename}"
            upload_bytes(data, object_name, file.content_type or "application/octet-stream")
            new_version = doc["current_version"] + 1
            uploader_name = resolve_user_name(cur, user_id)

            cur.execute(
                """UPDATE documents SET
                       object_name=%s, content_type=%s, size_bytes=%s, current_version=%s, updated_at=NOW()
                   WHERE id=%s RETURNING *""",
                (object_name, file.content_type, len(data), new_version, document_id),
            )
            updated = cur.fetchone()

            cur.execute(
                """INSERT INTO document_versions
                       (document_id, version_number, object_name, content_type, size_bytes, uploaded_by, uploaded_by_name, notes, company_id)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                (document_id, new_version, object_name, file.content_type, len(data), user_id, uploader_name, notes, company_id),
            )
            conn.commit()
            return attach_tags_and_url(cur, updated)


@app.get("/documents/{document_id}/versions/{version_number}/download")
def download_version(document_id: int, version_number: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT folder_id FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, doc["folder_id"], "view")

            cur.execute(
                "SELECT object_name FROM document_versions WHERE document_id = %s AND version_number = %s",
                (document_id, version_number),
            )
            version = cur.fetchone()
            if not version:
                raise HTTPException(status_code=404, detail="Versão não encontrada")
            return {"url": get_presigned_url(version["object_name"])}


@app.post("/documents/{document_id}/versions/{version_number}/restore")
def restore_version(document_id: int, version_number: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT * FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, doc["folder_id"], "edit")

            cur.execute(
                "SELECT * FROM document_versions WHERE document_id = %s AND version_number = %s",
                (document_id, version_number),
            )
            target = cur.fetchone()
            if not target:
                raise HTTPException(status_code=404, detail="Versão não encontrada")

            new_version = doc["current_version"] + 1
            uploader_name = resolve_user_name(cur, user_id)

            cur.execute(
                """UPDATE documents SET
                       object_name=%s, content_type=%s, size_bytes=%s, current_version=%s, updated_at=NOW()
                   WHERE id=%s RETURNING *""",
                (target["object_name"], target["content_type"], target["size_bytes"], new_version, document_id),
            )
            updated = cur.fetchone()

            cur.execute(
                """INSERT INTO document_versions
                       (document_id, version_number, object_name, content_type, size_bytes, uploaded_by, uploaded_by_name, notes, company_id)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
                (document_id, new_version, target["object_name"], target["content_type"], target["size_bytes"],
                 user_id, uploader_name, f"Restaurado a partir da versão {version_number}", company_id),
            )
            conn.commit()
            return attach_tags_and_url(cur, updated)


# ---------------------------------------------------------------------------
# Workflow de aprovação (upload → revisão → aprovação → publicado / rejeitado)
# ---------------------------------------------------------------------------

WORKFLOW_TRANSITIONS = {
    "submit":  {"from": ("enviado", "rejeitado"), "to": "em_revisao"},
    "approve": {"from": ("em_revisao",), "to": "aprovado"},
    "reject":  {"from": ("em_revisao",), "to": "rejeitado"},
    "publish": {"from": ("aprovado",), "to": "publicado"},
}


def run_workflow_transition(cur, document_id, user_id, company_id, action, notes):
    cur.execute("SELECT * FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
    doc = cur.fetchone()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento não encontrado")
    require_permission(cur, user_id, doc["folder_id"], "edit")

    rule = WORKFLOW_TRANSITIONS[action]
    if doc["workflow_state"] not in rule["from"]:
        raise HTTPException(
            status_code=409,
            detail=f"Transição inválida: documento está em '{doc['workflow_state']}'",
        )

    to_state = rule["to"]
    user_name = resolve_user_name(cur, user_id)

    cur.execute(
        "UPDATE documents SET workflow_state=%s, updated_at=NOW() WHERE id=%s AND company_id=%s RETURNING *",
        (to_state, document_id, company_id),
    )
    updated = cur.fetchone()

    cur.execute(
        """INSERT INTO document_workflow_log (document_id, from_state, to_state, user_id, user_name, notes, company_id)
           VALUES (%s, %s, %s, %s, %s, %s, %s)""",
        (document_id, doc["workflow_state"], to_state, user_id, user_name, notes, company_id),
    )

    if to_state == "aprovado":
        notify(cur, doc["owner_id"], "documento_aprovado", f"Documento aprovado: {doc['name']}", company_id, document_id)
    elif to_state == "rejeitado":
        notify(cur, doc["owner_id"], "documento_rejeitado", f"Documento rejeitado: {doc['name']}", company_id, document_id)

    return updated


@app.post("/documents/{document_id}/workflow/submit")
def workflow_submit(document_id: int, body: WorkflowTransition, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            updated = run_workflow_transition(cur, document_id, user_id, company_id, "submit", body.notes)
            conn.commit()
            return attach_tags_and_url(cur, updated)


@app.post("/documents/{document_id}/workflow/approve")
def workflow_approve(document_id: int, body: WorkflowTransition, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            updated = run_workflow_transition(cur, document_id, user_id, company_id, "approve", body.notes)
            conn.commit()
            return attach_tags_and_url(cur, updated)


@app.post("/documents/{document_id}/workflow/reject")
def workflow_reject(document_id: int, body: WorkflowTransition, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            updated = run_workflow_transition(cur, document_id, user_id, company_id, "reject", body.notes)
            conn.commit()
            return attach_tags_and_url(cur, updated)


@app.post("/documents/{document_id}/workflow/publish")
def workflow_publish(document_id: int, body: WorkflowTransition, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            updated = run_workflow_transition(cur, document_id, user_id, company_id, "publish", body.notes)
            conn.commit()
            return attach_tags_and_url(cur, updated)


@app.get("/documents/{document_id}/workflow/history")
def workflow_history(document_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT folder_id FROM documents WHERE id = %s AND company_id = %s", (document_id, company_id))
            doc = cur.fetchone()
            if not doc:
                raise HTTPException(status_code=404, detail="Documento não encontrado")
            require_permission(cur, user_id, doc["folder_id"], "view")

            cur.execute(
                "SELECT * FROM document_workflow_log WHERE document_id = %s ORDER BY created_at DESC",
                (document_id,),
            )
            return cur.fetchall()


# ---------------------------------------------------------------------------
# Tags
# ---------------------------------------------------------------------------

@app.get("/tags")
def list_tags(user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            ensure_default_tags(cur, company_id)
            conn.commit()
            cur.execute("SELECT * FROM doc_tags WHERE company_id = %s ORDER BY name", (company_id,))
            return cur.fetchall()


@app.post("/tags", status_code=201)
def create_tag(body: TagCreate, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT 1 FROM doc_tags WHERE name = %s AND company_id = %s", (body.name, company_id))
            if cur.fetchone():
                raise HTTPException(status_code=409, detail="Já existe uma etiqueta com este nome")
            # A verificação acima não é atómica — dois pedidos concorrentes podem passar
            # ambos por ela; o índice único (name, company_id) na BD é o backstop real.
            try:
                cur.execute(
                    "INSERT INTO doc_tags (name, color, company_id) VALUES (%s, %s, %s) RETURNING *",
                    (body.name, body.color, company_id),
                )
            except psycopg2.errors.UniqueViolation:
                conn.rollback()
                raise HTTPException(status_code=409, detail="Já existe uma etiqueta com este nome")
            conn.commit()
            return cur.fetchone()


# ---------------------------------------------------------------------------
# Permissões (por empresa / departamento / cargo / utilizador, herdadas na árvore de pastas)
# ---------------------------------------------------------------------------

VALID_SCOPE_TYPES = {"company", "department", "role", "user"}


@app.get("/permissions")
def list_permissions(folder_id: int | None = None, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            if folder_id is not None:
                cur.execute(
                    "SELECT * FROM document_permissions WHERE folder_id = %s AND company_id = %s ORDER BY id",
                    (folder_id, company_id),
                )
            else:
                cur.execute(
                    "SELECT * FROM document_permissions WHERE company_id = %s ORDER BY folder_id, id",
                    (company_id,),
                )
            return cur.fetchall()


@app.post("/permissions", status_code=201)
def create_permission(body: PermissionCreate, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    if body.scope_type not in VALID_SCOPE_TYPES:
        raise HTTPException(status_code=422, detail="scope_type inválido")
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT 1 FROM folders WHERE id = %s AND company_id = %s", (body.folder_id, company_id))
            if not cur.fetchone():
                raise HTTPException(status_code=404, detail="Pasta não encontrada")
            cur.execute(
                """INSERT INTO document_permissions
                       (scope_type, scope_value, folder_id, can_view, can_edit, can_delete, company_id)
                   VALUES (%s, %s, %s, %s, %s, %s, %s) RETURNING *""",
                (body.scope_type, body.scope_value, body.folder_id, body.can_view, body.can_edit, body.can_delete, company_id),
            )
            conn.commit()
            return cur.fetchone()


@app.delete("/permissions/{permission_id}", status_code=204)
def delete_permission(permission_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM document_permissions WHERE id = %s AND company_id = %s", (permission_id, company_id))
            conn.commit()


@app.get("/permissions/effective")
def effective_permission(folder_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT 1 FROM folders WHERE id = %s AND company_id = %s", (folder_id, company_id))
            if not cur.fetchone():
                raise HTTPException(status_code=404, detail="Pasta não encontrada")
            return resolve_folder_permission(cur, user_id, folder_id)


# ---------------------------------------------------------------------------
# Modelos (biblioteca de templates, com preenchimento automático de campos)
# ---------------------------------------------------------------------------

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@app.get("/templates")
def list_templates(category: str | None = None, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            if category:
                cur.execute(
                    "SELECT * FROM doc_templates WHERE category = %s AND company_id = %s ORDER BY name",
                    (category, company_id),
                )
            else:
                cur.execute("SELECT * FROM doc_templates WHERE company_id = %s ORDER BY name", (company_id,))
            return cur.fetchall()


@app.post("/templates", status_code=201)
def create_template(
    name: str,
    category: str | None = None,
    file: UploadFile = File(...),
    user_id: str = Depends(get_current_user),
    company_id: int = Depends(get_current_company),
):
    data = file.file.read()
    object_name = f"templates/{uuid.uuid4()}_{file.filename}"
    upload_bytes(data, object_name, file.content_type or "application/octet-stream")

    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            creator_name = resolve_user_name(cur, user_id)
            cur.execute(
                """INSERT INTO doc_templates
                       (name, category, object_name, content_type, size_bytes, created_by, created_by_name, company_id)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s) RETURNING *""",
                (name, category, object_name, file.content_type, len(data), user_id, creator_name, company_id),
            )
            conn.commit()
            return cur.fetchone()


@app.delete("/templates/{template_id}", status_code=204)
def delete_template(template_id: int, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT * FROM doc_templates WHERE id = %s AND company_id = %s", (template_id, company_id))
            template = cur.fetchone()
            if not template:
                raise HTTPException(status_code=404, detail="Modelo não encontrado")
            delete_object(template["object_name"])
            cur.execute("DELETE FROM doc_templates WHERE id = %s AND company_id = %s", (template_id, company_id))
            conn.commit()


def fill_docx_placeholders(data: bytes, fields: dict[str, str]) -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    for paragraph in doc.paragraphs:
        for key, value in fields.items():
            placeholder = f"{{{{{key}}}}}"
            if placeholder in paragraph.text:
                for run in paragraph.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, value)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@app.post("/templates/{template_id}/use", status_code=201)
def use_template(template_id: int, body: TemplateUse, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT * FROM doc_templates WHERE id = %s AND company_id = %s", (template_id, company_id))
            template = cur.fetchone()
            if not template:
                raise HTTPException(status_code=404, detail="Modelo não encontrado")
            cur.execute("SELECT 1 FROM folders WHERE id = %s AND company_id = %s", (body.folder_id, company_id))
            if not cur.fetchone():
                raise HTTPException(status_code=404, detail="Pasta não encontrada")
            require_permission(cur, user_id, body.folder_id, "edit")

            response = get_minio_client().get_object(BUCKET_NAME, template["object_name"])
            data = response.read()

            substituted = False
            if template["content_type"] == DOCX_MIME and body.fields:
                data = fill_docx_placeholders(data, body.fields)
                substituted = True

            name = body.document_name or template["name"]
            object_name = f"{body.folder_id}/{uuid.uuid4()}_{name}"
            upload_bytes(data, object_name, template["content_type"] or "application/octet-stream")

            owner_name = resolve_user_name(cur, user_id)
            cur.execute(
                """INSERT INTO documents
                       (folder_id, name, object_name, content_type, size_bytes,
                        category, owner_id, owner_name, workflow_state, company_id)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'enviado', %s)
                   RETURNING *""",
                (body.folder_id, name, object_name, template["content_type"], len(data),
                 template["category"], user_id, owner_name, company_id),
            )
            doc = cur.fetchone()
            cur.execute(
                """INSERT INTO document_versions
                       (document_id, version_number, object_name, content_type, size_bytes, uploaded_by, uploaded_by_name, company_id)
                   VALUES (%s, 1, %s, %s, %s, %s, %s, %s)""",
                (doc["id"], object_name, template["content_type"], len(data), user_id, owner_name, company_id),
            )
            conn.commit()
            result = attach_tags_and_url(cur, doc)
            result["substituted"] = substituted
            return result


# ---------------------------------------------------------------------------
# Notificações
# ---------------------------------------------------------------------------

def generate_expiry_notifications(cur, user_id, company_id):
    cur.execute("""
        SELECT id, name, owner_id, expiry_date FROM documents
        WHERE deleted_at IS NULL AND owner_id = %s AND company_id = %s AND expiry_date IS NOT NULL
              AND expiry_date <= CURRENT_DATE + INTERVAL '30 days'
    """, (user_id, company_id))
    for doc in cur.fetchall():
        notif_type = "documento_expirado" if doc["expiry_date"] < date.today() else "documento_a_expirar"
        cur.execute(
            "SELECT 1 FROM doc_notifications WHERE document_id = %s AND type = %s",
            (doc["id"], notif_type),
        )
        if not cur.fetchone():
            message = (
                f"Documento expirado: {doc['name']}" if notif_type == "documento_expirado"
                else f"Documento a expirar em breve: {doc['name']}"
            )
            notify(cur, user_id, notif_type, message, company_id, doc["id"])


@app.get("/notifications")
def list_notifications(unread_only: bool = False, user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            generate_expiry_notifications(cur, user_id, company_id)
            conn.commit()

            clauses = ["user_id = %s"]
            params = [user_id]
            if unread_only:
                clauses.append("is_read = false")
            cur.execute(
                f"SELECT * FROM doc_notifications WHERE {' AND '.join(clauses)} ORDER BY created_at DESC LIMIT 100",
                params,
            )
            return cur.fetchall()


@app.patch("/notifications/{notification_id}/read")
def mark_notification_read(notification_id: int, user_id: str = Depends(get_current_user)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute(
                "UPDATE doc_notifications SET is_read = true WHERE id = %s AND user_id = %s RETURNING *",
                (notification_id, user_id),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="Notificação não encontrada")
            conn.commit()
            return row


@app.post("/notifications/read-all")
def mark_all_notifications_read(user_id: str = Depends(get_current_user)):
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("UPDATE doc_notifications SET is_read = true WHERE user_id = %s", (user_id,))
            conn.commit()


# ---------------------------------------------------------------------------
# Stats (dashboard)
# ---------------------------------------------------------------------------

@app.get("/stats")
def get_stats(user_id: str = Depends(get_current_user), company_id: int = Depends(get_current_company)):
    with get_connection() as conn:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT COUNT(*) AS n FROM documents WHERE company_id = %s AND deleted_at IS NULL", (company_id,))
            total = cur.fetchone()["n"]

            cur.execute("""
                SELECT COUNT(*) AS n FROM documents
                WHERE company_id = %s AND deleted_at IS NULL AND expiry_date IS NOT NULL AND expiry_date < CURRENT_DATE
            """, (company_id,))
            expired = cur.fetchone()["n"]

            cur.execute("""
                SELECT COUNT(*) AS n FROM documents
                WHERE company_id = %s AND deleted_at IS NULL AND expiry_date IS NOT NULL
                      AND expiry_date >= CURRENT_DATE AND expiry_date <= CURRENT_DATE + INTERVAL '30 days'
            """, (company_id,))
            expiring_soon = cur.fetchone()["n"]

            cur.execute("""
                SELECT COUNT(*) AS n FROM documents
                WHERE company_id = %s AND deleted_at IS NULL AND signature_status IN ('pendente', 'em_assinatura')
            """, (company_id,))
            pending_signature = cur.fetchone()["n"]

            cur.execute(
                "SELECT COUNT(*) AS n FROM documents WHERE company_id = %s AND deleted_at IS NULL AND status = 'arquivado'",
                (company_id,),
            )
            archived = cur.fetchone()["n"]

            cur.execute("""
                SELECT COUNT(*) AS n FROM documents
                WHERE company_id = %s AND deleted_at IS NULL AND created_at >= date_trunc('month', NOW())
            """, (company_id,))
            uploads_this_month = cur.fetchone()["n"]

            cur.execute(
                "SELECT COALESCE(SUM(size_bytes), 0) AS n FROM documents WHERE company_id = %s AND deleted_at IS NULL",
                (company_id,),
            )
            storage_used_bytes = cur.fetchone()["n"]

    return {
        "total": total,
        "expired": expired,
        "expiring_soon": expiring_soon,
        "pending_signature": pending_signature,
        "archived": archived,
        "uploads_this_month": uploads_this_month,
        "storage_used_bytes": storage_used_bytes,
    }
